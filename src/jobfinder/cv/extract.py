"""Plain-text extraction from CV files."""

from __future__ import annotations

import re
from pathlib import Path

import docx
import pdfplumber
from pdfminer.pdfparser import PDFSyntaxError
from pdfplumber.utils.exceptions import PdfminerException

MIN_TEXT_CHARS = 200
_BLANK_LINES = re.compile(r"\n{3,}")
_TRAILING_SPACE = re.compile(r"[ \t]+\n")


class CVError(Exception):
    """The CV cannot be used (unsupported, unreadable or without extractable text)."""


def _pdf_text(path: Path) -> str:
    try:
        with pdfplumber.open(path) as pdf:
            return "\n".join(page.extract_text() or "" for page in pdf.pages)
    except (PdfminerException, PDFSyntaxError, ValueError) as exc:
        raise CVError(f"{path.name}: not a readable PDF ({exc})") from exc


def _docx_text(path: Path) -> str:
    try:
        document = docx.Document(str(path))
    except Exception as exc:  # python-docx raises zipfile/KeyError/ValueError variants
        raise CVError(f"{path.name}: not a readable DOCX ({type(exc).__name__})") from exc
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def normalise_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = _TRAILING_SPACE.sub("\n", text)
    return _BLANK_LINES.sub("\n\n", text).strip()


def extract_text(path: Path) -> str:
    """Extract text from a PDF or DOCX CV; raise ``CVError`` if there is too little."""
    if not path.is_file():
        raise CVError(f"{path}: file not found")
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        raw = _pdf_text(path)
    elif suffix == ".docx":
        raw = _docx_text(path)
    else:
        raise CVError(f"{path.name}: unsupported format {suffix or '(none)'}; use PDF or DOCX")
    text = normalise_text(raw)
    if len(text) < MIN_TEXT_CHARS:
        raise CVError(
            f"{path.name}: only {len(text)} characters of text could be extracted. "
            "Scanned or image-only CVs are not supported; export a text-based PDF or DOCX."
        )
    return text
