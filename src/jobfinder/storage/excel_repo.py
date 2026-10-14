"""OpportunityRepository backed by an Excel workbook (opportunities.xlsx).

The workbook is loaded into memory when the repository is opened; ``flush`` writes it
back. Every write:

1. takes an exclusive fcntl lock (``<workbook>.lock``),
2. re-reads the user-owned columns (application_status, notes) from the workbook on
   disk and merges them in, so edits made while jobfinder was running are kept,
3. copies the current workbook into ``rotated/`` (the newest 7 copies are kept),
4. writes a temporary file in the same directory and ``os.replace``s it into place,
   so a crash leaves the previous workbook intact.

Datetimes are stored as UTC without a timezone (Excel has none); list columns are
joined with "; ". Only the known columns are kept: extra columns added by hand are not
preserved.
"""

from __future__ import annotations

import fcntl
import os
import shutil
from collections.abc import Callable, Iterator
from contextlib import contextmanager
from datetime import UTC, datetime
from enum import Enum
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from jobfinder.models import (
    OPPORTUNITY_COLUMNS,
    PROFILE_COLUMNS,
    RUN_COLUMNS,
    SIGHTING_COLUMNS,
    USER_OWNED_FIELDS,
    Decision,
    ListingStatus,
    Opportunity,
    ProfileRecord,
    RunStats,
    Sighting,
    UserFields,
    profile_rows,
    utcnow,
)
from jobfinder.pipeline.normalise import norm_company
from jobfinder.storage.tables import LIST_COLUMNS

OPPORTUNITIES_SHEET = "Opportunities"
OTHER_JOBS_SHEET = "Other jobs"  # hidden: every job not shown on Opportunities
JOB_SHEETS = (OPPORTUNITIES_SHEET, OTHER_JOBS_SHEET)
SIGHTINGS_SHEET = "Sightings"
RUNS_SHEET = "Runs"
PROFILE_SHEET = "Profile"
ROTATED_COPIES = 7
MAX_RUN_ROWS = 2000  # about three months of hourly runs
LIST_SEPARATOR = "; "
MAX_CELL_CHARS = 32_767
_FORMULA_PREFIXES = ("=", "+", "-", "@")
_DATETIME_FORMAT = "yyyy-mm-dd hh:mm"
_WIDE_COLUMNS = {"title": 40, "company": 24, "summary": 60, "score_reason": 50, "notes": 40}


class ExcelRepositoryError(Exception):
    pass


# --------------------------------------------------------------------------- cell codec


def _to_cell(name: str, value: Any) -> Any:
    if value is None:
        return None
    if name in LIST_COLUMNS:
        value = LIST_SEPARATOR.join(value)
        return value or None
    if isinstance(value, Enum):
        return value.value
    if isinstance(value, datetime):
        return value.astimezone(UTC).replace(tzinfo=None)
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", value)[:MAX_CELL_CHARS]
    return value


def _from_cell(name: str, value: Any) -> Any:
    if name in LIST_COLUMNS:
        if value in (None, ""):
            return []
        return [part.strip() for part in str(value).split(";") if part.strip()]
    if value == "":
        return None
    if isinstance(value, datetime):
        return value.replace(tzinfo=UTC) if value.tzinfo is None else value.astimezone(UTC)
    if name in ("relevance_score", "prefilter_score", "profile_version") and value is not None:
        return int(value)
    if name in USER_OWNED_FIELDS and value is not None:
        return str(value)
    return value


def _write_rows(sheet: Worksheet, columns: tuple[str, ...], rows: list[dict[str, Any]]) -> None:
    sheet.append(list(columns))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row_index, row in enumerate(rows, start=2):
        for col_index, name in enumerate(columns, start=1):
            value = row.get(name)
            if value is None:
                continue
            cell = sheet.cell(row=row_index, column=col_index, value=value)
            if isinstance(value, str) and value.startswith(_FORMULA_PREFIXES):
                cell.data_type = "s"  # never let fetched text become a formula
            elif isinstance(value, datetime):
                cell.number_format = _DATETIME_FORMAT
    sheet.freeze_panes = "A2"
    if rows:
        sheet.auto_filter.ref = sheet.dimensions
    for col_index, name in enumerate(columns, start=1):
        sheet.column_dimensions[get_column_letter(col_index)].width = _WIDE_COLUMNS.get(name, 16)


def _read_sheet(sheet: Worksheet) -> list[dict[str, Any]]:
    rows = sheet.iter_rows(values_only=True)
    header = next(rows, None)
    if header is None:
        return []
    names = [str(h) if h is not None else "" for h in header]
    result = []
    for values in rows:
        if all(v is None for v in values):
            continue
        result.append({name: value for name, value in zip(names, values, strict=False) if name})
    return result


# --------------------------------------------------------------------------- repository


class ExcelRepository:
    def __init__(
        self,
        path: Path,
        *,
        rotated_dir: Path | None = None,
        clock: Callable[[], datetime] = utcnow,
        read_only: bool = False,
        visible_decisions: frozenset[Decision] | None = None,
    ) -> None:
        self.path = path
        self.read_only = read_only
        # None: every job on the Opportunities sheet. Otherwise only active jobs with these
        # decisions (or with your own notes) are shown; the rest go to a hidden sheet.
        self.visible_decisions = visible_decisions
        self.rotated_dir = rotated_dir or path.parent / "rotated"
        self._lock_path = path.with_name(path.name + ".lock")
        self._clock = clock
        self._opportunities: dict[str, Opportunity] = {}
        self._sightings: list[Sighting] = []
        self._runs: list[dict[str, Any]] = []
        self._profile: list[tuple[str, str]] = []
        self._dirty = False
        path.parent.mkdir(parents=True, exist_ok=True)
        with self._locked():
            self._load()

    # ------------------------------------------------------------------ io

    @contextmanager
    def _locked(self) -> Iterator[None]:
        fd = os.open(self._lock_path, os.O_RDWR | os.O_CREAT, 0o600)
        try:
            fcntl.flock(fd, fcntl.LOCK_EX)
            yield
        finally:
            fcntl.flock(fd, fcntl.LOCK_UN)
            os.close(fd)

    def _open_workbook(self) -> Workbook | None:
        if not self.path.exists():
            return None
        try:
            return load_workbook(self.path, read_only=True, data_only=True)
        except Exception as exc:  # openpyxl raises many types for damaged files
            raise ExcelRepositoryError(
                f"{self.path} cannot be read ({type(exc).__name__}: {exc}); restore a copy "
                f"from {self.rotated_dir}"
            ) from exc

    def _load(self) -> None:
        workbook = self._open_workbook()
        if workbook is None:
            return
        try:
            for sheet_name in JOB_SHEETS:
                if sheet_name not in workbook.sheetnames:
                    continue
                for row in _read_sheet(workbook[sheet_name]):
                    data = {n: _from_cell(n, row.get(n)) for n in OPPORTUNITY_COLUMNS}
                    opportunity = Opportunity.model_validate(data)
                    self._opportunities[opportunity.opportunity_id] = opportunity
            if SIGHTINGS_SHEET in workbook.sheetnames:
                for row in _read_sheet(workbook[SIGHTINGS_SHEET]):
                    data = {n: _from_cell(n, row.get(n)) for n in SIGHTING_COLUMNS}
                    self._sightings.append(Sighting.model_validate(data))
            if RUNS_SHEET in workbook.sheetnames:
                self._runs = [
                    {k: _from_cell(k, v) for k, v in row.items()}
                    for row in _read_sheet(workbook[RUNS_SHEET])
                ]
            if PROFILE_SHEET in workbook.sheetnames:
                self._profile = [
                    (str(row.get("field") or ""), str(row.get("value") or ""))
                    for row in _read_sheet(workbook[PROFILE_SHEET])
                    if row.get("field")
                ]
        finally:
            workbook.close()

    def _disk_user_fields(self) -> dict[str, UserFields]:
        workbook = self._open_workbook()
        if workbook is None:
            return {}
        try:
            return {
                str(row["opportunity_id"]): UserFields(
                    application_status=_from_cell(
                        "application_status", row.get("application_status")
                    ),
                    notes=_from_cell("notes", row.get("notes")),
                )
                for sheet_name in JOB_SHEETS
                if sheet_name in workbook.sheetnames
                for row in _read_sheet(workbook[sheet_name])
                if row.get("opportunity_id")
            }
        finally:
            workbook.close()

    def is_visible(self, opportunity: Opportunity) -> bool:
        if self.visible_decisions is None:
            return True
        tracked = bool(opportunity.application_status or opportunity.notes)
        return opportunity.decision in self.visible_decisions and (
            opportunity.listing_status is ListingStatus.ACTIVE or tracked
        )

    def _rotate(self) -> None:
        if not self.path.exists():
            return
        self.rotated_dir.mkdir(parents=True, exist_ok=True)
        stamp = self._clock().astimezone(UTC).strftime("%Y%m%dT%H%M%S%fZ")
        shutil.copy2(self.path, self.rotated_dir / f"{self.path.stem}-{stamp}{self.path.suffix}")
        copies = sorted(self.rotated_dir.glob(f"{self.path.stem}-*{self.path.suffix}"))
        for old in copies[:-ROTATED_COPIES]:
            old.unlink()

    def _build_workbook(self) -> Workbook:
        workbook = Workbook()
        default = workbook.active
        if default is not None:
            workbook.remove(default)
        opportunities = sorted(self._opportunities.values(), key=lambda o: o.discovered_at)
        shown = [o for o in opportunities if self.is_visible(o)]
        shown.sort(key=lambda o: -(o.relevance_score or 0))  # best first
        _write_rows(
            workbook.create_sheet(OPPORTUNITIES_SHEET),
            OPPORTUNITY_COLUMNS,
            [{n: _to_cell(n, getattr(o, n)) for n in OPPORTUNITY_COLUMNS} for o in shown],
        )
        _write_rows(
            workbook.create_sheet(SIGHTINGS_SHEET),
            SIGHTING_COLUMNS,
            [{n: _to_cell(n, getattr(s, n)) for n in SIGHTING_COLUMNS} for s in self._sightings],
        )
        _write_rows(
            workbook.create_sheet(RUNS_SHEET),
            RUN_COLUMNS,
            [{k: _to_cell(k, v) for k, v in row.items()} for row in self._runs[-MAX_RUN_ROWS:]],
        )
        _write_rows(
            workbook.create_sheet(PROFILE_SHEET),
            PROFILE_COLUMNS,
            [{"field": f, "value": _to_cell("value", v)} for f, v in self._profile],
        )
        if self.visible_decisions is not None:
            other = workbook.create_sheet(OTHER_JOBS_SHEET)
            _write_rows(
                other,
                OPPORTUNITY_COLUMNS,
                [
                    {n: _to_cell(n, getattr(o, n)) for n in OPPORTUNITY_COLUMNS}
                    for o in opportunities
                    if not self.is_visible(o)
                ],
            )
            other.sheet_state = "hidden"
        return workbook

    def flush(self) -> None:
        """Merge user edits from disk, then atomically replace the workbook."""
        if self.read_only:
            return
        with self._locked():
            for opportunity_id, fields in self._disk_user_fields().items():
                current = self._opportunities.get(opportunity_id)
                if current is not None:
                    self._opportunities[opportunity_id] = current.model_copy(
                        update=fields.model_dump()
                    )
            if not self._dirty and self.path.exists():
                return
            workbook = self._build_workbook()
            tmp = self.path.with_name(f".{self.path.name}.{os.getpid()}.tmp")
            try:
                workbook.save(tmp)
                with tmp.open("rb") as handle:
                    os.fsync(handle.fileno())
                tmp.chmod(0o600)
                self._rotate()
                tmp.replace(self.path)
            except BaseException:
                tmp.unlink(missing_ok=True)
                raise
            self._dirty = False

    def close(self) -> None:
        self.flush()

    # ------------------------------------------------------------------ queries

    def _sighting_owner(self, predicate: Callable[[Sighting], bool]) -> Opportunity | None:
        for sighting in self._sightings:
            if predicate(sighting):
                return self._opportunities.get(sighting.opportunity_id)
        return None

    def _earliest(self, predicate: Callable[[Opportunity], bool]) -> Opportunity | None:
        matches = [o for o in self._opportunities.values() if predicate(o)]
        return min(matches, key=lambda o: o.discovered_at) if matches else None

    def get_by_source_key(self, source: str, source_job_id: str) -> Opportunity | None:
        found = self._earliest(lambda o: o.source == source and o.source_job_id == source_job_id)
        return found or self._sighting_owner(
            lambda s: s.source == source and s.source_job_id == source_job_id
        )

    def get_by_canonical_url(self, url: str) -> Opportunity | None:
        found = self._earliest(lambda o: o.canonical_url == url)
        return found or self._sighting_owner(lambda s: s.canonical_url == url)

    def get_by_fingerprint(self, fingerprint: str) -> Opportunity | None:
        return self._earliest(lambda o: o.fingerprint == fingerprint)

    def fuzzy_candidates(self, company_norm: str) -> list[Opportunity]:
        return [o for o in self._opportunities.values() if norm_company(o.company) == company_norm]

    def upsert(self, opportunity: Opportunity) -> None:
        existing = self._opportunities.get(opportunity.opportunity_id)
        if existing is not None:
            opportunity = opportunity.model_copy(
                update={f: getattr(existing, f) for f in USER_OWNED_FIELDS}
            )
        self._opportunities[opportunity.opportunity_id] = opportunity
        self._dirty = True

    def add_sighting(self, sighting: Sighting) -> None:
        if any(
            s.source == sighting.source and s.source_job_id == sighting.source_job_id
            for s in self._sightings
        ):
            return
        self._sightings.append(sighting)
        self._dirty = True

    def mark_status(self, opportunity_id: str, status: ListingStatus) -> None:
        current = self._opportunities.get(opportunity_id)
        if current is not None and current.listing_status is not status:
            self._opportunities[opportunity_id] = current.model_copy(
                update={"listing_status": status}
            )
            self._dirty = True

    def iter_active(self) -> Iterator[Opportunity]:
        active = [
            o for o in self._opportunities.values() if o.listing_status is ListingStatus.ACTIVE
        ]
        yield from sorted(active, key=lambda o: o.discovered_at)

    def iter_all(self) -> Iterator[Opportunity]:
        yield from sorted(self._opportunities.values(), key=lambda o: o.discovered_at)

    def user_fields(self, opportunity_id: str) -> UserFields | None:
        current = self._opportunities.get(opportunity_id)
        if current is None:
            return None
        return UserFields(application_status=current.application_status, notes=current.notes)

    def sightings_for(self, opportunity_id: str) -> list[Sighting]:
        return [s for s in self._sightings if s.opportunity_id == opportunity_id]

    def count(self) -> int:
        return len(self._opportunities)

    def has_source(self, source: str) -> bool:
        return any(o.source == source for o in self._opportunities.values()) or any(
            s.source == source for s in self._sightings
        )

    def append_run(self, stats: RunStats) -> None:
        self._runs.append(stats.row())
        self._dirty = True

    def write_profile(self, record: ProfileRecord) -> None:
        rows = profile_rows(record)
        if rows != self._profile:
            self._profile = rows
            self._dirty = True
