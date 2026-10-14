"""Contract tests: both backends must behave identically."""

from __future__ import annotations

from collections.abc import Iterator
from datetime import UTC, datetime, timedelta
from pathlib import Path

import pytest

from jobfinder.models import (
    Decision,
    EmploymentType,
    ListingStatus,
    MatchRule,
    Opportunity,
    RemoteType,
    Seniority,
    Sighting,
    UserFields,
)
from jobfinder.storage.excel_repo import ExcelRepository
from jobfinder.storage.repository import OpportunityRepository
from jobfinder.storage.sql_repo import SqlRepository

T0 = datetime(2026, 10, 1, 9, 30, 15, 123456, tzinfo=UTC)


def make(opportunity_id: str = "01J0000000000000000000000A", **kw: object) -> Opportunity:
    data: dict[str, object] = {
        "opportunity_id": opportunity_id,
        "fingerprint": "f" * 40,
        "canonical_url": "https://jobs.lever.co/acme/1",
        "source": "lever",
        "source_job_id": "1",
        "title": "Backend Engineer",
        "company": "Acme GmbH",
        "location_city": "Berlin",
        "location_country": "DE",
        "remote_type": RemoteType.HYBRID,
        "employment_type": EmploymentType.FULL_TIME,
        "seniority": Seniority.JUNIOR,
        "required_skills": ["Python", "SQL"],
        "posting_url": "https://jobs.lever.co/acme/1",
        "apply_url": "https://jobs.lever.co/acme/1/apply",
        "published_at": T0 - timedelta(days=2),
        "discovered_at": T0,
        "last_seen_at": T0,
        "decision": Decision.PENDING_AI,
    }
    data.update(kw)
    return Opportunity.model_validate(data)


@pytest.fixture(params=["sqlite", "excel"])
def backend(request: pytest.FixtureRequest) -> str:
    return str(request.param)


def open_repo(backend: str, tmp_path: Path) -> OpportunityRepository:
    if backend == "sqlite":
        return SqlRepository.sqlite(tmp_path / "opportunities.db")
    return ExcelRepository(tmp_path / "opportunities.xlsx")


@pytest.fixture
def repo(backend: str, tmp_path: Path) -> Iterator[OpportunityRepository]:
    r = open_repo(backend, tmp_path)
    yield r
    r.close()


def test_roundtrip_through_reopen(backend: str, tmp_path: Path) -> None:
    original = make(red_flags=["keyword stuffing"], relevance_score=77, notes=None)
    repo = open_repo(backend, tmp_path)
    repo.upsert(original)
    repo.close()
    reopened = open_repo(backend, tmp_path)
    loaded = reopened.get_by_source_key("lever", "1")
    reopened.close()
    assert loaded == original
    assert loaded is not None
    assert loaded.discovered_at.tzinfo is UTC


def test_lookups(repo: OpportunityRepository) -> None:
    opp = make()
    repo.upsert(opp)
    assert repo.get_by_source_key("lever", "1") == opp
    assert repo.get_by_source_key("lever", "2") is None
    assert repo.get_by_source_key("greenhouse", "1") is None
    assert repo.get_by_canonical_url("https://jobs.lever.co/acme/1") == opp
    assert repo.get_by_fingerprint("f" * 40) == opp
    assert repo.fuzzy_candidates("acme") == [opp]
    assert repo.fuzzy_candidates("globex") == []
    assert repo.count() == 1


def test_lookups_through_sightings(repo: OpportunityRepository) -> None:
    opp = make()
    repo.upsert(opp)
    repo.add_sighting(
        Sighting(
            opportunity_id=opp.opportunity_id,
            source="remotive",
            source_job_id="99",
            posting_url="https://remotive.com/remote-jobs/x-99",
            canonical_url="https://remotive.com/remote-jobs/x-99",
            seen_at=T0,
            match_rule=MatchRule.FINGERPRINT,
        )
    )
    assert repo.get_by_source_key("remotive", "99") == opp
    assert repo.get_by_canonical_url("https://remotive.com/remote-jobs/x-99") == opp
    assert len(repo.sightings_for(opp.opportunity_id)) == 1


def test_duplicate_sighting_ignored(repo: OpportunityRepository) -> None:
    opp = make()
    repo.upsert(opp)
    sighting = Sighting(
        opportunity_id=opp.opportunity_id,
        source="remotive",
        source_job_id="99",
        posting_url="https://remotive.com/x",
        canonical_url="https://remotive.com/x",
        seen_at=T0,
        match_rule=MatchRule.CANONICAL_URL,
    )
    repo.add_sighting(sighting)
    repo.add_sighting(sighting.model_copy(update={"seen_at": T0 + timedelta(hours=1)}))
    assert len(repo.sightings_for(opp.opportunity_id)) == 1


def test_upsert_never_overwrites_user_fields(backend: str, tmp_path: Path) -> None:
    repo = open_repo(backend, tmp_path)
    repo.upsert(make())
    if isinstance(repo, SqlRepository):
        from sqlalchemy import update

        from jobfinder.storage.tables import opportunities

        with repo._engine.begin() as conn:
            conn.execute(
                update(opportunities).values(application_status="applied", notes="call Fri")
            )
    else:
        repo.close()
        _edit_excel_user_fields(tmp_path / "opportunities.xlsx", "applied", "call Fri")
        repo = open_repo(backend, tmp_path)
    repo.upsert(make(title="Backend Engineer II", application_status=None, notes="overwritten?"))
    assert repo.user_fields("01J0000000000000000000000A") == UserFields(
        application_status="applied", notes="call Fri"
    )
    found = repo.get_by_source_key("lever", "1")
    assert found is not None
    assert found.title == "Backend Engineer II"
    repo.close()


def test_mark_status_and_iter_active(repo: OpportunityRepository) -> None:
    a = make("01J0000000000000000000000A")
    b = make(
        "01J0000000000000000000000B", source_job_id="2", discovered_at=T0 + timedelta(minutes=1)
    )
    repo.upsert(b)
    repo.upsert(a)
    assert [o.opportunity_id for o in repo.iter_active()] == [a.opportunity_id, b.opportunity_id]
    repo.mark_status(a.opportunity_id, ListingStatus.EXPIRED)
    assert [o.opportunity_id for o in repo.iter_active()] == [b.opportunity_id]
    found = repo.get_by_source_key("lever", "1")
    assert found is not None
    assert found.listing_status is ListingStatus.EXPIRED


def test_user_fields_unknown_id(repo: OpportunityRepository) -> None:
    assert repo.user_fields("nope") is None


def _edit_excel_user_fields(path: Path, status: str, notes: str) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    sheet = workbook["Opportunities"]
    header = [c.value for c in sheet[1]]
    sheet.cell(row=2, column=header.index("application_status") + 1).value = status
    sheet.cell(row=2, column=header.index("notes") + 1).value = notes
    workbook.save(path)


def test_has_source_and_run_rows(backend: str, tmp_path: Path) -> None:
    from jobfinder.models import RunStats
    from tests.helpers import make_profile_record

    repo = open_repo(backend, tmp_path)
    assert not repo.has_source("lever")
    repo.upsert(make())
    assert repo.has_source("lever")
    repo.add_sighting(
        Sighting(
            opportunity_id=make().opportunity_id,
            source="remotive",
            source_job_id="9",
            posting_url="https://r.io/9",
            canonical_url="https://r.io/9",
            seen_at=T0,
            match_rule=MatchRule.FINGERPRINT,
        )
    )
    assert repo.has_source("remotive")
    repo.append_run(RunStats(run_id="r1", started_at=T0, status="success"))
    repo.write_profile(make_profile_record())
    repo.write_profile(make_profile_record(version=2))
    repo.close()


def test_sqlite_visible_view(tmp_path: Path) -> None:
    from sqlalchemy import text

    repo = SqlRepository.sqlite(tmp_path / "o.db", frozenset({Decision.STRONG}))
    repo.upsert(make("01J0000000000000000000000A", source_job_id="1", decision=Decision.STRONG))
    repo.upsert(make("01J0000000000000000000000B", source_job_id="2", decision=Decision.FILTERED))
    with repo._engine.connect() as conn:
        rows = conn.execute(text("SELECT source_job_id FROM visible_opportunities")).all()
    assert rows == [("1",)]
    repo.close()
    # Changing the setting replaces the view.
    repo = SqlRepository.sqlite(tmp_path / "o.db", frozenset({Decision.STRONG, Decision.FILTERED}))
    with repo._engine.connect() as conn:
        assert len(conn.execute(text("SELECT * FROM visible_opportunities")).all()) == 2
    repo.close()
