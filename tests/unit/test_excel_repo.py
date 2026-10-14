from __future__ import annotations

from datetime import UTC, datetime, timedelta
from pathlib import Path

import pytest
from openpyxl import load_workbook

from jobfinder.models import OPPORTUNITY_COLUMNS, UserFields
from jobfinder.storage.excel_repo import (
    ROTATED_COPIES,
    ExcelRepository,
    ExcelRepositoryError,
)
from tests.unit.test_repositories import T0, make


def _set_user_fields(path: Path, row: int, status: str | None, notes: str | None) -> None:
    workbook = load_workbook(path)
    sheet = workbook["Opportunities"]
    header = [c.value for c in sheet[1]]
    # ws.cell(..., value=None) leaves the cell unchanged, so assign .value directly.
    sheet.cell(row=row, column=header.index("application_status") + 1).value = status
    sheet.cell(row=row, column=header.index("notes") + 1).value = notes
    workbook.save(path)


def test_layout(tmp_path: Path) -> None:
    path = tmp_path / "opportunities.xlsx"
    repo = ExcelRepository(path)
    repo.upsert(make())
    repo.flush()
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Opportunities", "Sightings", "Runs", "Profile"]
    sheet = workbook["Opportunities"]
    assert tuple(c.value for c in sheet[1]) == OPPORTUNITY_COLUMNS
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == sheet.dimensions
    assert path.stat().st_mode & 0o777 == 0o600


def test_user_edits_made_during_a_run_are_kept(tmp_path: Path) -> None:
    path = tmp_path / "opportunities.xlsx"
    repo = ExcelRepository(path)
    repo.upsert(make())
    repo.flush()
    # The user edits the workbook while jobfinder still has it loaded...
    _set_user_fields(path, 2, "interview", "2nd round Tue")
    # ...and jobfinder then updates the same row and writes.
    repo.upsert(make(title="Backend Engineer (updated)", last_seen_at=T0 + timedelta(hours=1)))
    repo.flush()
    reopened = ExcelRepository(path)
    assert reopened.user_fields(make().opportunity_id) == UserFields(
        application_status="interview", notes="2nd round Tue"
    )
    found = reopened.get_by_source_key("lever", "1")
    assert found is not None
    assert found.title == "Backend Engineer (updated)"


def test_user_can_clear_their_fields(tmp_path: Path) -> None:
    path = tmp_path / "opportunities.xlsx"
    repo = ExcelRepository(path)
    repo.upsert(make(application_status="applied"))
    repo.flush()
    _set_user_fields(path, 2, None, None)
    repo.upsert(make(application_status="applied"))
    repo.flush()
    assert ExcelRepository(path).user_fields(make().opportunity_id) == UserFields()


def test_failed_write_keeps_previous_workbook(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "opportunities.xlsx"
    repo = ExcelRepository(path)
    repo.upsert(make())
    repo.flush()
    before = path.read_bytes()

    def crash(self: object, filename: object) -> None:
        Path(str(filename)).write_bytes(b"half a workbook")
        raise OSError("killed mid-write")

    monkeypatch.setattr("openpyxl.Workbook.save", crash)
    repo.upsert(make("01J0000000000000000000000B", source_job_id="2"))
    with pytest.raises(OSError, match="killed mid-write"):
        repo.flush()
    monkeypatch.undo()
    assert path.read_bytes() == before
    assert ExcelRepository(path).count() == 1
    assert not list(tmp_path.glob(".*.tmp"))


def test_rotation_keeps_seven_copies(tmp_path: Path) -> None:
    path = tmp_path / "opportunities.xlsx"
    ticks = iter(T0 + timedelta(minutes=i) for i in range(100))
    repo = ExcelRepository(path, clock=lambda: next(ticks))
    for i in range(10):
        repo.upsert(make(f"01J00000000000000000000{i:03d}", source_job_id=str(i)))
        repo.flush()
    copies = sorted((tmp_path / "rotated").iterdir())
    assert len(copies) == ROTATED_COPIES
    # The newest copy is the workbook as it was before the last write.
    assert ExcelRepository(copies[-1]).count() == 9


def test_no_write_when_nothing_changed(tmp_path: Path) -> None:
    path = tmp_path / "opportunities.xlsx"
    repo = ExcelRepository(path)
    repo.upsert(make())
    repo.flush()
    mtime = path.stat().st_mtime_ns
    repo.flush()
    assert path.stat().st_mtime_ns == mtime


@pytest.mark.parametrize("payload", ['=HYPERLINK("http://evil")', "+1+1", "-2+3", "@SUM(A1)"])
def test_formula_injection_stored_as_text(tmp_path: Path, payload: str) -> None:
    path = tmp_path / "opportunities.xlsx"
    repo = ExcelRepository(path)
    repo.upsert(make(title=payload))
    repo.flush()
    sheet = load_workbook(path)["Opportunities"]
    cell = sheet.cell(row=2, column=OPPORTUNITY_COLUMNS.index("title") + 1)
    assert cell.data_type == "s"
    assert cell.value == payload
    found = ExcelRepository(path).get_by_source_key("lever", "1")
    assert found is not None
    assert found.title == payload


def test_illegal_characters_removed(tmp_path: Path) -> None:
    path = tmp_path / "opportunities.xlsx"
    repo = ExcelRepository(path)
    repo.upsert(make(title="Engineer\x00\x0b II"))
    repo.flush()
    found = ExcelRepository(path).get_by_source_key("lever", "1")
    assert found is not None
    assert found.title == "Engineer II"


def test_runs_and_profile_sheets(tmp_path: Path) -> None:
    from jobfinder.models import RUN_COLUMNS, RunStats
    from tests.helpers import make_profile_record

    path = tmp_path / "opportunities.xlsx"
    repo = ExcelRepository(path)
    repo.append_run(RunStats(run_id="r1", started_at=T0, status="success", new=3))
    repo.write_profile(make_profile_record())
    repo.flush()
    reopened = ExcelRepository(path)
    reopened.append_run(RunStats(run_id="r2", started_at=T0, status="partial"))
    reopened.flush()
    workbook = load_workbook(path)
    runs = list(workbook["Runs"].iter_rows(values_only=True))
    assert runs[0] == RUN_COLUMNS
    assert [(r[0], r[3]) for r in runs[1:]] == [("r1", "success"), ("r2", "partial")]
    profile = dict(workbook["Profile"].iter_rows(min_row=2, values_only=True))
    assert profile["profile_version"] == "1"
    assert profile["seniority"] == "junior"


def test_datetimes_stored_as_utc(tmp_path: Path) -> None:
    from datetime import timezone

    path = tmp_path / "opportunities.xlsx"
    local = datetime(2026, 10, 1, 11, 30, tzinfo=timezone(timedelta(hours=2)))
    repo = ExcelRepository(path)
    repo.upsert(make(published_at=local))
    repo.flush()
    sheet = load_workbook(path)["Opportunities"]
    cell = sheet.cell(row=2, column=OPPORTUNITY_COLUMNS.index("published_at") + 1)
    assert cell.value == datetime(2026, 10, 1, 9, 30)  # noqa: DTZ001 - Excel has no zones
    found = ExcelRepository(path).get_by_source_key("lever", "1")
    assert found is not None
    assert found.published_at == local
    assert found.published_at is not None and found.published_at.tzinfo is UTC


def test_corrupt_workbook_explains_recovery(tmp_path: Path) -> None:
    path = tmp_path / "opportunities.xlsx"
    path.write_bytes(b"not a workbook")
    with pytest.raises(ExcelRepositoryError, match="restore a copy"):
        ExcelRepository(path)


def test_only_chosen_decisions_are_shown(tmp_path: Path) -> None:
    from jobfinder.models import Decision, ListingStatus
    from jobfinder.storage.excel_repo import OTHER_JOBS_SHEET

    path = tmp_path / "opportunities.xlsx"
    repo = ExcelRepository(path, visible_decisions=frozenset({Decision.STRONG}))
    repo.upsert(
        make(
            "01J0000000000000000000000A",
            source_job_id="1",
            decision=Decision.STRONG,
            relevance_score=80,
        )
    )
    repo.upsert(
        make(
            "01J0000000000000000000000B",
            source_job_id="2",
            decision=Decision.STRONG,
            relevance_score=95,
        )
    )
    repo.upsert(make("01J0000000000000000000000C", source_job_id="3", decision=Decision.FILTERED))
    repo.upsert(
        make(
            "01J0000000000000000000000D",
            source_job_id="4",
            decision=Decision.REJECT,
            relevance_score=30,
        )
    )
    repo.upsert(
        make(
            "01J0000000000000000000000E",
            source_job_id="5",
            decision=Decision.STRONG,
            listing_status=ListingStatus.EXPIRED,
        )
    )
    repo.upsert(
        make(
            "01J0000000000000000000000F",
            source_job_id="6",
            decision=Decision.STRONG,
            listing_status=ListingStatus.EXPIRED,
            application_status="applied",
        )
    )
    repo.flush()

    workbook = load_workbook(path)
    shown = [
        r[OPPORTUNITY_COLUMNS.index("source_job_id")]
        for r in workbook["Opportunities"].iter_rows(min_row=2, values_only=True)
    ]
    assert shown == ["2", "1", "6"]  # strong, active or tracked by you; best score first
    other = workbook[OTHER_JOBS_SHEET]
    assert other.sheet_state == "hidden"
    assert sorted(
        r[OPPORTUNITY_COLUMNS.index("source_job_id")]
        for r in other.iter_rows(min_row=2, values_only=True)
    ) == ["3", "4", "5"]

    # Everything is still known after a reload (dedupe needs it), and edits made on the
    # hidden sheet are kept like edits on the visible one.
    other_header = [c.value for c in other[1]]
    other.cell(row=2, column=other_header.index("notes") + 1).value = "looked at it"
    workbook.save(path)
    reopened = ExcelRepository(path, visible_decisions=frozenset({Decision.STRONG}))
    assert reopened.count() == 6
    hidden_id = next(o.opportunity_id for o in reopened.iter_all() if o.notes == "looked at it")
    reopened.upsert(reopened._opportunities[hidden_id].model_copy(update={"notes": None}))
    reopened.flush()
    assert ExcelRepository(path).user_fields(hidden_id).notes == "looked at it"  # type: ignore[union-attr]


def test_no_hidden_sheet_without_a_selection(tmp_path: Path) -> None:
    path = tmp_path / "opportunities.xlsx"
    repo = ExcelRepository(path)
    repo.upsert(make())
    repo.flush()
    assert "Other jobs" not in load_workbook(path).sheetnames
