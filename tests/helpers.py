"""Test helpers: CV file builders and a scripted HTTP transport for the openai SDK."""

from __future__ import annotations

import json
from collections import deque
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import docx
import httpx2

FIXTURES = Path(__file__).resolve().parent / "fixtures"


def _pdf_escape(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def make_pdf(path: Path, lines: list[str]) -> Path:
    """Write a minimal one-page, text-based PDF (Helvetica, WinAnsi) with ``lines``."""
    ops = " ".join(f"({_pdf_escape(line)}) '" for line in lines)
    content = f"BT /F1 9 Tf 40 810 Td 11 TL {ops} ET".encode("cp1252")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n".encode()
    out += f"startxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(bytes(out))
    return path


def make_docx(path: Path, paragraphs: list[str], table: list[list[str]] | None = None) -> Path:
    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        t = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                t.cell(r, c).text = value
    document.save(str(path))
    return path


def sample_cv_lines() -> list[str]:
    return (FIXTURES / "cv" / "sample_cv.txt").read_text(encoding="utf-8").splitlines()


def load_ai_fixture(name: str) -> dict[str, Any]:
    data: dict[str, Any] = json.loads((FIXTURES / "ai" / name).read_text(encoding="utf-8"))
    return data


def completion_body(
    content: str, *, model: str = "openai/gpt-oss-120b", **usage: int
) -> dict[str, Any]:
    return {
        "id": "chatcmpl-test",
        "object": "chat.completion",
        "created": 1791900000,
        "model": model,
        "choices": [
            {
                "index": 0,
                "message": {"role": "assistant", "content": content},
                "finish_reason": "stop",
            }
        ],
        "usage": {
            "prompt_tokens": usage.get("prompt_tokens", 100),
            "completion_tokens": usage.get("completion_tokens", 50),
            "total_tokens": usage.get("prompt_tokens", 100) + usage.get("completion_tokens", 50),
        },
    }


@dataclass
class ScriptedTransport:
    """Serves queued responses in order and records every request."""

    responses: deque[httpx2.Response] = field(default_factory=deque)
    requests: list[httpx2.Request] = field(default_factory=list)

    def add(
        self, status: int = 200, body: Any = None, headers: dict[str, str] | None = None
    ) -> ScriptedTransport:
        self.responses.append(httpx2.Response(status, json=body, headers=headers or {}))
        return self

    def add_fixture(self, name: str) -> ScriptedTransport:
        fixture = load_ai_fixture(name)
        return self.add(fixture["status"], fixture["body"], fixture["headers"])

    def handler(self, request: httpx2.Request) -> httpx2.Response:
        self.requests.append(request)
        if not self.responses:
            raise AssertionError(f"unexpected request: {request.method} {request.url}")
        return self.responses.popleft()

    def client(self) -> httpx2.AsyncClient:
        return httpx2.AsyncClient(transport=httpx2.MockTransport(self.handler))

    def bodies(self) -> list[dict[str, Any]]:
        return [json.loads(r.content) for r in self.requests]


@dataclass
class FakeSleep:
    calls: list[float] = field(default_factory=list)

    async def __call__(self, seconds: float) -> None:
        self.calls.append(seconds)


def load_source_fixture(name: str) -> dict[str, Any]:
    """A recorded source response: {"recorded_at", "headers", "body"}."""
    data: dict[str, Any] = json.loads(
        (FIXTURES / "sources" / f"{name}.json").read_text(encoding="utf-8")
    )
    return data


@dataclass
class FakeClock:
    """A clock that only moves when something sleeps on it."""

    now: Any  # datetime
    sleeps: list[float] = field(default_factory=list)

    def __call__(self) -> Any:
        return self.now

    async def sleep(self, seconds: float) -> None:
        from datetime import timedelta

        self.sleeps.append(seconds)
        self.now = self.now + timedelta(seconds=seconds)


def make_raw_job(**kw: Any) -> Any:
    """A RawJob that fits PROFILE_EXAMPLE-like profiles (junior Python backend)."""
    from datetime import UTC, datetime

    from jobfinder.models import RawJob

    data: dict[str, Any] = {
        "source": "lever",
        "source_job_id": "job-1",
        "is_ats": True,
        "title": "Junior Backend Developer (Python)",
        "company": "Logistico",
        "locations": ["Paris, France"],
        "remote_hint": "hybrid",
        "employment_type_hint": "Full-time",
        "description_text": (
            "About the role\nYou will build REST APIs in Python with Django and PostgreSQL "
            "for our logistics platform. 1+ years of experience with Python.\n\n"
            "Benefits\nFree lunch, gym membership.\n\n"
            "We are an equal opportunity employer."
        ),
        "posting_url": "https://jobs.lever.co/logistico/job-1",
        "apply_url": "https://jobs.lever.co/logistico/job-1/apply",
        "published_at": datetime(2026, 10, 12, tzinfo=UTC),
    }
    data.update(kw)
    return RawJob.model_validate(data)


def make_profile_record(version: int = 1) -> Any:
    from datetime import UTC, datetime

    from jobfinder.ai.examples import PROFILE_EXAMPLE
    from jobfinder.models import ProfileRecord

    return ProfileRecord(
        version=version,
        created_at=datetime(2026, 10, 1, tzinfo=UTC),
        cv_sha256="a" * 64,
        cv_filename="cv.pdf",
        overrides_sha256="b" * 64,
        ai_provider="groq",
        ai_model="openai/gpt-oss-120b",
        extracted=PROFILE_EXAMPLE,
        profile=PROFILE_EXAMPLE,
    )
