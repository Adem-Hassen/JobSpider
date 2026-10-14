"""``jobfinder run`` end to end: fake sources, recorded AI replies, mocked Telegram."""

from __future__ import annotations

import asyncio
import json
from datetime import UTC, datetime, timedelta
from typing import Any

import httpx2
import pytest
import respx
from openpyxl import load_workbook

from jobfinder.ai.client import OpenAICompatibleClient
from jobfinder.config import AppConfig, SourceName, load_config
from jobfinder.cv.profile_store import ProfileStore
from jobfinder.models import Decision, RawJob
from jobfinder.pipeline import run as run_module
from jobfinder.pipeline.run import RunDeps, execute_run
from jobfinder.sources.base import FetchContext
from jobfinder.storage.excel_repo import ExcelRepository
from jobfinder.storage.state_store import SourceState, StateStore
from tests.helpers import FakeSleep, load_ai_fixture, make_profile_record, make_raw_job

NOW = datetime(2026, 10, 14, 12, tzinfo=UTC)
TOKEN = "999:telegram-test-token"
TELEGRAM = f"https://api.telegram.org/bot{TOKEN}/sendMessage"
HEALTHCHECK = "https://hc-ping.com/abc-123"
SALES = {
    "source_job_id": "job-2",
    "title": "Account Executive",
    "description_text": "Sell our platform.",
    "posting_url": "https://jobs.lever.co/logistico/job-2",
    "apply_url": "https://jobs.lever.co/logistico/job-2/apply",
}


class FakeSource:
    def __init__(
        self,
        name: str,
        jobs: list[RawJob] | None = None,
        error: Exception | None = None,
        delay: float = 0,
    ) -> None:
        self.name = name
        self.jobs = jobs or []
        self.error = error
        self.delay = delay

    async def fetch(self, ctx: FetchContext, state: SourceState) -> list[RawJob]:
        if self.delay:
            await asyncio.sleep(self.delay)
        if self.error:
            raise self.error
        return self.jobs


def repeating_ai_client(config: AppConfig) -> OpenAICompatibleClient:
    fixture = load_ai_fixture("groq_score_response.json")

    def handler(request: httpx2.Request) -> httpx2.Response:
        return httpx2.Response(200, json=fixture["body"], headers=fixture["headers"])

    return OpenAICompatibleClient(
        name="groq",
        provider=config.provider,
        api_key="k",
        user_agent="t",
        http_client=httpx2.AsyncClient(transport=httpx2.MockTransport(handler)),
        sleep=FakeSleep(),
    )


@pytest.fixture
def config(base_env: dict[str, str], monkeypatch: pytest.MonkeyPatch, data_dir: Any) -> AppConfig:
    monkeypatch.setenv("TELEGRAM_BOT_TOKEN", TOKEN)
    monkeypatch.setenv("TELEGRAM_CHAT_ID", "42")
    monkeypatch.setenv("HEALTHCHECK_URL", HEALTHCHECK)
    cfg = load_config(env_file=None)
    ProfileStore(cfg.profiles_dir).save(make_profile_record())
    return cfg


def deps(config: AppConfig, sources: dict[SourceName, Any], now: datetime = NOW) -> RunDeps:
    return RunDeps(
        clock=lambda: now, sleep=FakeSleep(), sources=sources, ai_client=repeating_ai_client(config)
    )


def mock_outside_world(telegram_status: int = 200) -> tuple[respx.Route, respx.Route]:
    respx.get(url__regex=r"https://[^/]+/robots\.txt").respond(404)
    respx.head(url__startswith="https://jobs.lever.co/").respond(200)
    telegram = respx.post(TELEGRAM).respond(
        telegram_status,
        json={"ok": telegram_status == 200, "description": "Bad Request: chat not found"},
    )
    healthcheck = respx.get(HEALTHCHECK).respond(200)
    return telegram, healthcheck


def lever(*jobs: RawJob) -> dict[SourceName, Any]:
    return {SourceName.LEVER: FakeSource("lever", list(jobs))}


@respx.mock
async def test_full_run(config: AppConfig) -> None:
    telegram, healthcheck = mock_outside_world()
    outcome = await execute_run(config, deps(config, lever(make_raw_job(), make_raw_job(**SALES))))
    stats = outcome.stats
    assert stats.status == "success", stats.errors
    assert (stats.new, stats.filtered, stats.scored, stats.notified, stats.links_resolved) == (
        2,
        1,
        1,
        1,
        1,
    )
    assert stats.sources["lever"].jobs_fetched == 2
    assert stats.ai_cost_usd > 0 and stats.peak_rss_mb > 0

    message = json.loads(telegram.calls[0].request.content)["text"]
    assert "Junior Backend Developer (Python)" in message and "Listing on Lever" in message
    assert healthcheck.call_count == 1

    workbook = load_workbook(config.data_dir / "opportunities.xlsx")
    runs = list(workbook["Runs"].iter_rows(values_only=True))
    assert runs[1][0] == stats.run_id and runs[1][3] == "success"
    assert (
        dict(workbook["Profile"].iter_rows(min_row=2, values_only=True))["profile_version"] == "1"
    )
    repo = ExcelRepository(config.data_dir / "opportunities.xlsx", read_only=True)
    strong = repo.get_by_source_key("lever", "job-1")
    assert strong is not None and strong.decision is Decision.STRONG
    assert strong.apply_url_final == "https://jobs.lever.co/logistico/job-1/apply"
    with StateStore(config.state_db_path) as store:
        assert store.recent_runs(1)[0].status == "success"
        assert store.get_source_state("lever").last_success_at == NOW


@respx.mock
async def test_second_run_adds_no_duplicates_and_sends_nothing_again(config: AppConfig) -> None:
    telegram, _ = mock_outside_world()
    jobs = lever(make_raw_job(), make_raw_job(**SALES))
    await execute_run(config, deps(config, jobs))
    later = NOW + timedelta(hours=7)  # lever's interval is 6 h
    stats = (await execute_run(config, deps(config, jobs, later))).stats
    assert (stats.new, stats.updated, stats.notified) == (0, 2, 0)
    assert telegram.call_count == 1
    assert ExcelRepository(config.data_dir / "opportunities.xlsx", read_only=True).count() == 2


@respx.mock
async def test_failing_source_never_aborts_the_others(config: AppConfig) -> None:
    mock_outside_world()
    sources = {
        **lever(make_raw_job()),
        SourceName.GREENHOUSE: FakeSource("greenhouse", error=RuntimeError("board exploded")),
    }
    stats = (await execute_run(config, deps(config, sources))).stats
    assert stats.status == "partial"
    assert stats.sources["greenhouse"].status == "failed"
    assert any("board exploded" in e for e in stats.errors)
    assert stats.new == 1 and stats.notified == 1
    assert respx.get(HEALTHCHECK).call_count == 1  # partial runs still ping


@respx.mock
async def test_telegram_failure_is_retried_next_run(config: AppConfig) -> None:
    telegram, _ = mock_outside_world(telegram_status=400)
    first = (await execute_run(config, deps(config, lever(make_raw_job())))).stats
    assert first.status == "partial" and first.notified == 0
    assert any("chat not found" in e for e in first.errors)
    assert all(TOKEN not in e for e in first.errors)
    telegram.respond(200, json={"ok": True})
    second = (await execute_run(config, deps(config, {}, NOW + timedelta(minutes=10)))).stats
    assert second.notified == 1


@respx.mock
async def test_unwritable_storage_fails_the_run(config: AppConfig) -> None:
    _, healthcheck = mock_outside_world()
    (config.data_dir / "opportunities.xlsx").write_bytes(b"corrupt")
    stats = (await execute_run(config, deps(config, lever(make_raw_job())))).stats
    assert stats.status == "failed"
    assert any("restore a copy" in e for e in stats.errors)
    assert healthcheck.call_count == 0  # the dead-man switch will notice
    with StateStore(config.state_db_path) as store:
        assert store.recent_runs(1)[0].status == "failed"


@respx.mock
async def test_slow_source_hits_the_fetch_deadline(
    config: AppConfig, monkeypatch: pytest.MonkeyPatch
) -> None:
    mock_outside_world()
    monkeypatch.setattr(run_module, "FETCH_DEADLINE", timedelta(seconds=0.05))
    sources = {**lever(make_raw_job()), SourceName.ASHBY: FakeSource("ashby", delay=5)}
    stats = (await execute_run(config, deps(config, sources))).stats
    assert stats.sources["ashby"].status == "failed"
    assert stats.new == 1
    with StateStore(config.state_db_path) as store:
        assert store.get_source_state("ashby").consecutive_failures == 1


@respx.mock
async def test_without_profile_jobs_wait_for_scoring(config: AppConfig) -> None:
    mock_outside_world()
    for path in config.profiles_dir.iterdir():
        path.unlink()
    stats = (await execute_run(config, deps(config, lever(make_raw_job())))).stats
    assert stats.status == "partial"
    assert stats.pending_ai == 1
    assert any("no profile yet" in e for e in stats.errors)


@respx.mock
async def test_jobs_in_other_countries_are_not_stored(config: AppConfig) -> None:
    mock_outside_world()
    us = make_raw_job(
        source_job_id="us",
        locations=["Austin, TX"],
        posting_url="https://jobs.lever.co/logistico/us",
    )
    us_remote = make_raw_job(
        source_job_id="usr",
        locations=["Remote, United States"],
        remote_hint="remote",
        posting_url="https://jobs.lever.co/logistico/usr",
    )
    stats = (await execute_run(config, deps(config, lever(make_raw_job(), us, us_remote)))).stats
    assert (stats.new, stats.out_of_scope) == (1, 2)
    repo = ExcelRepository(config.data_dir / "opportunities.xlsx", read_only=True)
    assert repo.count() == 1 and repo.get_by_source_key("lever", "us") is None


@respx.mock
async def test_store_other_countries_keeps_them_as_filtered(
    base_env: dict[str, str], monkeypatch: pytest.MonkeyPatch, edit_yaml: Any
) -> None:
    mock_outside_world()
    edit_yaml("settings.yaml", lambda d: d.update(store_other_countries=True))
    cfg = load_config(env_file=None)
    ProfileStore(cfg.profiles_dir).save(make_profile_record())
    us = make_raw_job(
        source_job_id="us",
        locations=["Austin, TX"],
        posting_url="https://jobs.lever.co/logistico/us",
    )
    stats = (await execute_run(cfg, deps(cfg, lever(us)))).stats
    assert (stats.new, stats.out_of_scope, stats.filtered) == (1, 0, 1)
    stored = ExcelRepository(cfg.data_dir / "opportunities.xlsx", read_only=True).get_by_source_key(
        "lever", "us"
    )
    assert stored is not None and stored.decision is Decision.FILTERED
    assert stored.score_reason is not None and "located in US" in stored.score_reason
